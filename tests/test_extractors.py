import asyncio

import pandas as pd
from PIL import Image
from docx import Document

from pipelines.ingestion.extractors.csv_extractor import extract_csv
from pipelines.ingestion.extractors.docx_extractor import extract_docx
from pipelines.ingestion.extractors.image_extractor import extract_image
from pipelines.ingestion.extractors.url_extractor import extract_url


def test_docx_extractor_paragraphs_tables_and_headings(tmp_path):
	doc = Document()
	doc.add_heading("Main Heading", level=1)
	doc.add_paragraph("Paragraph text")
	table = doc.add_table(rows=2, cols=2)
	table.rows[0].cells[0].text = "A"
	table.rows[0].cells[1].text = "B"
	table.rows[1].cells[0].text = "1"
	table.rows[1].cells[1].text = "2"
	path = tmp_path / "sample.docx"
	doc.save(path)

	pages = extract_docx(str(path))
	assert any(page.content_type == "heading" for page in pages)
	assert any(page.content_type == "table" for page in pages)
	assert any("heading_path" in page.metadata for page in pages if page.metadata)


def test_csv_extractor_small_and_large_modes(tmp_path):
	small = pd.DataFrame({"col": list(range(10))})
	small_path = tmp_path / "small.csv"
	small.to_csv(small_path, index=False)
	small_pages = extract_csv(str(small_path))
	assert len(small_pages) == 1
	assert "|" in small_pages[0].content

	large = pd.DataFrame({"num": list(range(1100)), "cat": ["x"] * 1100})
	large_path = tmp_path / "large.csv"
	large.to_csv(large_path, index=False)
	large_pages = extract_csv(str(large_path))
	assert len(large_pages) == 11
	assert "CSV statistical summary" in large_pages[0].content


def test_image_extractor_combines_description_and_ocr(tmp_path, monkeypatch):
	img_path = tmp_path / "img.png"
	image = Image.new("RGB", (2000, 1200), color="white")
	image.save(img_path)

	monkeypatch.setattr(
		"pipelines.ingestion.extractors.image_extractor.pytesseract.image_to_string",
		lambda _img: "detected text",
	)
	pages = extract_image(str(img_path))
	assert len(pages) == 1
	assert "Text found in image: detected text" in pages[0].content
	assert pages[0].metadata["width"] <= 1024
	assert pages[0].metadata["height"] <= 1024


def test_url_extractor_metadata_and_content(monkeypatch):
	class FakeResponse:
		status_code = 200
		text = "<html><head><title>T</title></head><body>Body</body></html>"
		url = "https://example.com/article"

		def raise_for_status(self):
			return None

	class FakeAsyncClient:
		def __init__(self, *args, **kwargs):
			pass

		async def __aenter__(self):
			return self

		async def __aexit__(self, exc_type, exc, tb):
			return False

		async def get(self, _url):
			return FakeResponse()

	class FakeMetadata:
		title = "Example"
		author = "Author"
		date = "2026-01-01"
		url = "https://example.com/article"

	monkeypatch.setattr(
		"pipelines.ingestion.extractors.url_extractor._robots_allowed",
		lambda _url: asyncio.sleep(0, result=True),
	)
	monkeypatch.setattr(
		"pipelines.ingestion.extractors.url_extractor.httpx.AsyncClient",
		FakeAsyncClient,
	)
	monkeypatch.setattr(
		"pipelines.ingestion.extractors.url_extractor.trafilatura.extract",
		lambda *args, **kwargs: "Extracted body text",
	)
	monkeypatch.setattr(
		"pipelines.ingestion.extractors.url_extractor.trafilatura.extract_metadata",
		lambda _html: FakeMetadata(),
	)

	pages = asyncio.run(extract_url("https://example.com/article"))
	assert pages[0].content == "Extracted body text"
	assert pages[0].metadata["title"] == "Example"
