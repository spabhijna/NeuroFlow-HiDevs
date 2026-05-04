from __future__ import annotations

from docx import Document

from pipelines.ingestion import ExtractedPage


def _table_to_markdown(table) -> str:
	rows: list[str] = []
	for row in table.rows:
		cells = [cell.text.strip().replace("\n", " ") for cell in row.cells]
		rows.append("| " + " | ".join(cells) + " |")
	if not rows:
		return ""
	if len(rows) == 1:
		return rows[0]
	separator = "| " + " | ".join(["---"] * len(table.rows[0].cells)) + " |"
	return "\n".join([rows[0], separator, *rows[1:]])


def extract_docx(file_path: str) -> list[ExtractedPage]:
	document = Document(file_path)
	pages: list[ExtractedPage] = []
	heading_stack: dict[int, str] = {}
	page_number = 1

	for paragraph in document.paragraphs:
		text = paragraph.text.strip()
		if not text:
			continue

		style_name = paragraph.style.name if paragraph.style else ""
		metadata: dict[str, object] = {}
		content_type = "text"
		if style_name.startswith("Heading"):
			level_part = style_name.replace("Heading", "").strip()
			level = int(level_part) if level_part.isdigit() else 1
			heading_stack[level] = text
			for stale_level in [key for key in heading_stack if key > level]:
				del heading_stack[stale_level]
			hierarchy = [heading_stack[idx] for idx in sorted(heading_stack)]
			metadata = {
				"heading_level": level,
				"heading_path": hierarchy,
			}
			content_type = "heading"
		elif heading_stack:
			metadata = {
				"heading_path": [heading_stack[idx] for idx in sorted(heading_stack)],
			}

		pages.append(
			ExtractedPage(
				page_number=page_number,
				content=text,
				content_type=content_type,
				metadata=metadata,
			)
		)
		page_number += 1

	for table in document.tables:
		table_text = _table_to_markdown(table)
		if not table_text:
			continue
		metadata: dict[str, object] = {}
		if heading_stack:
			metadata["heading_path"] = [heading_stack[idx] for idx in sorted(heading_stack)]
		pages.append(
			ExtractedPage(
				page_number=page_number,
				content=table_text,
				content_type="table",
				metadata=metadata,
			)
		)
		page_number += 1

	return pages
